"""
Streamlit app: upload images + text, generate a polished .docx.

Run locally:   streamlit run app.py
Deploy:        push to GitHub -> connect repo at share.streamlit.io
"""

import io
from PIL import Image
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

st.set_page_config(page_title="Doc Builder", page_icon="📄")
st.title("📄 Document Builder")
st.caption("Fill in the details, upload images, download a Word doc.")

# ---------- Inputs ----------
title = st.text_input("Document title", "My Report")
intro = st.text_area("Intro / summary paragraph", "")

uploaded_files = st.file_uploader(
    "Upload images (in the order you want them to appear)",
    type=["png", "jpg", "jpeg"],
    accept_multiple_files=True,
)

captions = []
if uploaded_files:
    st.subheader("Captions (optional)")
    for f in uploaded_files:
        captions.append(st.text_input(f"Caption for {f.name}", key=f.name))

TARGET_WIDTH_IN = 6.0  # fits inside standard 1" margins on Letter/A4


def resized_image_stream(uploaded_file, max_width_in=TARGET_WIDTH_IN, dpi=150):
    """
    Open the uploaded image, downscale it if it's larger than needed
    (keeps the docx file size sane), and return (BytesIO, width_in, height_in)
    with the correct aspect ratio preserved.
    """
    img = Image.open(uploaded_file).convert("RGB")
    w_px, h_px = img.size
    aspect = h_px / w_px

    # Cap the pixel width to what's actually needed at the target DPI,
    # so a 6000px phone photo doesn't bloat the file.
    max_px_width = int(max_width_in * dpi)
    if w_px > max_px_width:
        new_w = max_px_width
        new_h = int(new_w * aspect)
        img = img.resize((new_w, new_h), Image.LANCZOS)

    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=88)
    buf.seek(0)

    width_in = max_width_in
    height_in = width_in * aspect
    return buf, width_in, height_in


def build_docx(title, intro, files, captions):
    doc = Document()

    # Base styling
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    heading = doc.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if intro.strip():
        doc.add_paragraph(intro.strip())

    for f, caption in zip(files, captions):
        f.seek(0)
        img_stream, width_in, height_in = resized_image_stream(f)

        # Keep every image within the page — cap height too (in case of
        # a very tall/narrow image) and re-derive width from that cap.
        MAX_HEIGHT_IN = 8.0
        if height_in > MAX_HEIGHT_IN:
            scale = MAX_HEIGHT_IN / height_in
            height_in = MAX_HEIGHT_IN
            width_in = width_in * scale

        doc.add_picture(img_stream, width=Inches(width_in))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if caption:
            cap = doc.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(9)

        doc.add_paragraph()  # spacer

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


if st.button("Generate document", type="primary"):
    if not uploaded_files and not intro.strip():
        st.warning("Add some text or at least one image first.")
    else:
        docx_bytes = build_docx(title, intro, uploaded_files or [], captions)
        st.success("Document generated.")
        st.download_button(
            "⬇️ Download .docx",
            data=docx_bytes,
            file_name=f"{title or 'document'}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
